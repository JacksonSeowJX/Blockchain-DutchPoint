"""Build matching final audit content without editing source section drafts."""
from pathlib import Path
import re, json, hashlib, zipfile
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
QA = ROOT / 'validation' / 'stage8'
QA.mkdir(exist_ok=True)
ASSETS = ROOT / 'diagrams' / 'report'
ASSETS.mkdir(exist_ok=True)
before = {str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in ROOT.glob('sections/*.md')}
(QA / 'preserved_sections.json').write_text(json.dumps(before, indent=2))

fontfile = Path('C:/Windows/Fonts/arial.ttf')
boldfile = Path('C:/Windows/Fonts/arialbd.ttf')
def diagram(name, rows, notes):
    """Simple vertical flow with labelled stages and adjacent explanatory notes."""
    w, h = 1700, 100 + len(rows)*215
    im = Image.new('RGB',(w,h),'white')
    d = ImageDraw.Draw(im)
    f = ImageFont.truetype(str(fontfile), 34)
    b = ImageFont.truetype(str(boldfile), 35)
    small = ImageFont.truetype(str(fontfile), 33)
    def wrapped(text, limit, font):
        lines=[]
        for para in text.split('\n'):
            line=''
            for word in para.split():
                if d.textlength(line+' '+word,font=font)>limit and line:
                    lines.append(line);line=word
                else: line=(line+' '+word).strip()
            lines.append(line)
        return lines
    for i,(title,body) in enumerate(rows):
        y=35+i*215
        d.rounded_rectangle((45,y,835,y+170),radius=12,fill='#edf2f6',outline='#64748b',width=2)
        d.text((68,y+15),title,font=b,fill='black')
        for j,line in enumerate(wrapped(body,744,f)):
            d.text((68,y+60+j*39),line,font=f,fill='black')
        if i<len(rows)-1:
            d.line((440,y+170,440,y+211),fill='#334155',width=4)
            d.polygon([(430,y+201),(450,y+201),(440,y+213)],fill='#334155')
        for j,line in enumerate(wrapped(notes[i],765,small)):
            d.text((880,y+10+j*38),line,font=small,fill='#242b35')
    im.save(ASSETS/name, dpi=(240,240))

diagram('figure_4_1.png',[
 ('Browser interface','React App, components and auctionMath'),
 ('Data and wallet hooks','useAuction reads and useWallet connectivity'),
 ('RPC access','Wallet BrowserProvider or fallback JsonRpcProvider'),
 ('Auction and token','DutchAuction and GavoraToken on the selected chain')], [
 'Vite serves static assets. Generated ABIs and deployment maps supply interfaces and addresses.',
 'Reads refresh state and contribution history. Display math derives estimates rather than authoritative settlement.',
 'Wallet authorizes writes. The fallback route provides reads. There is no backend or separate indexer.',
 'The auction reads balances, transfers claims and burns inventory through the token. Deployment scripts fund contracts and export metadata.'
 ])
diagram('figure_4_2.png',[
 ('Inventory and start','Issuer transfers tokens to the auction, owner starts'),
 ('Contribution','Bidder sends ETH, auction accepts up to remaining capacity'),
 ('Surplus handling','Auction sends surplus or records refund credit'),
 ('Settlement and distribution','After deadline, anyone may finalize'),
 ('Remainder','Anyone may burn after all participant claims')],[
 'Inventory is the auction address balance inside GavoraToken. The token has no owner or later mint function.',
 'Accepted ETH increases commitments. A crossing contribution can be only partly accepted.',
 'A credited bidder calls withdrawRefundTo for a chosen recipient. Credit withdrawal does not wait for finalization.',
 'Finalization burns unsold inventory. Bidders claim tokens and the current owner withdraws accepted proceeds independently.',
 'The auction burns its actual remaining token balance. Outstanding claims delay this burn, not issuer proceeds.'
 ])
diagram('figure_4_3.png',[
 ('Pending','Owner calls start after funding'),
 ('Active','Positive bids accepted while capacity remains'),
 ('Cleared','Crossing bid or elapsed time closes bidding'),
 ('AwaitingFinalization','Scheduled 20-minute deadline is reached'),
 ('Finalized','Anyone calls finalize, claims and proceeds become available')],[
 'Starting records the opening timestamp. Reading state does not submit a transaction.',
 'If the deadline arrives before clearing, Active goes directly to AwaitingFinalization.',
 'Time-only clearing is derived by views. Its stored fields and event are persisted by finalize.',
 'Deadline status takes priority over clearing. Clearing does not permit early finalization.',
 'Refund-credit withdrawals are independent of this lifecycle. Finalized remains the status during claims and later burns.'
 ])
diagram('figure_4_4.png',[
 ('Connect','User requests wallet access through useWallet'),
 ('Submit','App.run obtains a signer and requests contribute'),
 ('Confirm and mine','Wallet submits a signed transaction through RPC'),
 ('Receipt','App waits for the transaction receipt'),
 ('Refresh','useAuction reloads state and contribution history')],[
 'Account and chain are synchronized. A bundled deployment must exist for the selected chain.',
 'UI shows the wallet stage. Request rejection follows the error path through describeError.',
 'UI shows mining. DutchAuction executes the contribution and emits events if the call succeeds.',
 'A successful receipt leads to done and refresh. Reversion or receipt-wait failure leads to the failed UI stage.',
 'Fresh reads update the display. These are separate RPC calls, not an atomic snapshot or a guarantee of irreversible finality.'
 ])

titles = [
 'Initial Specification & System Scope',
 'Development Methodology & AI Tooling Infrastructure',
 'Specification Engineering & Prompt Logs',
 'AI-Generated Architecture',
 'Problems Discovered & Vulnerability Analysis',
 'Validation Process & Testing Strategy',
 'Improvements & Refinements',
 'Lessons Learned']

parts=[]
removals={
 1:['Source locators and claim classifications are retained'],
 2:['This section draws on the Stage 1 evidence register', 'The resulting commit sequence provides stronger evidence', 'Several benefits can be observed without inventing productivity', 'The limitations were equally important.'],
 3:['According to the edited record, our team requested a plan before implementation'],
 4:['The separation of responsibilities and smaller frontend modules plausibly', 'The four Mermaid sources are retained'],
 7:['The most substantial revision changed how the auction accepts a crossing bid.', 'Three implementation corrections accompanied this revision.'],
 8:[]}
for n,p in enumerate(sorted((ROOT/'sections').glob('0*.md')),1):
    s=p.read_text(encoding='utf-8')
    s=re.sub(r'^# .+\n',f'# {n} {titles[n-1]}\n',s,count=1)
    paras=s.split('\n\n')
    s='\n\n'.join(x for x in paras if not any(x.startswith(a) for a in removals.get(n,[])))
    if n==1:
        s=s.replace('Our application addresses', '## Purpose and users\n\nOur application addresses',1)
        s=s.replace('The original functional requirements were recorded', '## Requirements and scope boundaries\n\nThe original functional requirements were recorded',1)
        s=s.replace('The planned architecture separated', '## Initial architecture and economic policy\n\nThe planned architecture separated',1)
        s=s.replace('The expected transaction sequence reflected', '## Transactions and user interactions\n\nThe expected transaction sequence reflected',1)
        s=s.replace('The development log preserves a summary', '## Instructions to AI tools and later additions\n\nThe development log preserves a summary',1)
    if n==2:
        s=s.replace('The repository records a development process', '## Tooling and attribution\n\nThe repository records a development process',1)
        s=s.replace('The initial workflow was specification-led', '## Development and validation workflow\n\nThe initial workflow was specification-led',1)
        s=s.replace('Human involvement is documented', '## Human decisions and division of work\n\nHuman involvement is documented',1)
    if n==3:
        s=s.replace('This section completes Stage 3 without turning reconstructed history into original prompt evidence.','Reconstructed history is not presented as original prompt evidence.')
        s=s.replace('Neither the measurements nor the mutation were rerun in this stage.','Neither the measurements nor the mutation were rerun during this audit.')
    if n==4:
        s=s.replace('but are not a substitute for later vulnerability analysis.','and are evaluated in Section 5.')
        captions=['Overall application architecture','Contract and asset interactions','Auction status transitions','Wallet transaction flow']
        idx=[0]
        def fig(m):
            k=idx[0];idx[0]+=1
            return f'![Figure 4.{k+1} {captions[k]}](diagrams/report/figure_4_{k+1}.png)\n\nFigure 4.{k+1} {captions[k]}'
        s=re.sub(r'```mermaid\n.*?```',fig,s,flags=re.S)
    if n==5:
        s=s.replace('Stage 5 ends with this classification and evidence record. Stage 6 has not begun.','Section 6 records the subsequent validation and its limits.')
        s=s.replace('The detailed finding register in `../audit_findings.md`','The detailed finding register in `audit_findings.md` and the source map in Appendix B')
    if n==6:
        s=s.replace('## New execution during Stage 6','## New execution during the audit')
        s=s.replace('independently observed Stage 6 execution','independently observed audit execution')
    if n==7:
        s=s.replace('This is a confirmed policy revision,', 'The capped crossing-bid rule examined in Section 5 is a confirmed policy revision,',1)
        s=s.replace('| Improvement | Original limitation | Modification and outcome |','Table 7.1 summarizes the contract changes and their outcomes. The test-before-fix sequence is evaluated in Section 6.\n\n| Improvement | Original limitation | Modification and outcome |',1)
    s=s.replace('`../testing_results.md`','`testing_results.md`').replace('`../validation/`','`validation/`').replace('`../validation/stage6/`','`validation/stage6/`')
    s=s.replace('in this stage','during this audit').replace('during preparation of this section','during this audit')
    # Number subsections, including the original third-level headings.
    major=minor=0; out=[]
    for line in s.splitlines():
        if line.startswith('## '):
            major+=1;minor=0;line=f'## {n}.{major} '+line[3:]
        elif line.startswith('### '):
            minor+=1;line=f'### {n}.{major}.{minor} '+line[4:]
        out.append(line)
    parts.append('\n'.join(out).strip())

# Label all tables without changing their substantive cell text.
table_names={1:['Original functional requirements','Expected transactions in the initial design'],2:['AI tooling and conventional infrastructure'],3:['Representative prompt iterations and outcomes'],4:['Auction interface and permissions','Contract state groups','Frontend responsibilities','Architecture evolution'],5:['Historical findings and current risks'],6:['Existing test definitions'],7:['Contract refinements','Frontend and wallet refinements'],8:[]}
for n,s in enumerate(parts,1):
    count=0; lines=s.splitlines(); out=[]
    for i,line in enumerate(lines):
        if line.startswith('|') and (i==0 or not lines[i-1].startswith('|')):
            count+=1
            out.extend([f'Table {n}.{count} {table_names[n][count-1]}',''])
        out.append(line)
    parts[n-1]='\n'.join(out)

front='''# DutchPoint AI Audit Report

SC4053 Blockchain Technology

Nanyang Technological University

DutchPoint and the Gavora token

23 October 2026

Repository baseline `c36cb0cebe803db5ee138d89fe2e21e8cfb33ad5`

## Scope and evidence

This report evaluates our project's requirements, AI-assisted development, architecture, defects and validation. The implementation changed from pro-rata settlement to capped crossing bids, corrected lifecycle and proceeds behavior, and gained broader testing and a rebuilt interface. The audit also identifies five current source-verified frontend issues and four provisional risks. No unauthorized fund-extraction exploit was demonstrated in the current production contracts. [E43, E49–E62]

Confirmed findings are supported by inspected source, Git objects or the content of retained records. Inferred claims are interpretations that have not been reproduced. Unverified claims lack independent evidence. A historical test log confirms a retained result, not a fresh execution. We use first-person team language for activities documented in the repository, without assigning unsupported individual responsibilities.

The report compares initial commit `0a75a78` with the baseline above. Its conclusions are bounded by missing primary course documents and raw AI sessions, absent public deployment records and incomplete runtime validation. The audit did not modify application code or existing tests. Evidence identifiers refer to the repository source index in Appendix A and the companion evidence register.

## Table of contents

'''
for n,title in enumerate(titles,1):
    slug=re.sub(r'[^a-z0-9 -]','',f'{n} {title}'.lower()).replace(' ','-')
    front+=f'- [{n} {title}](#{slug})\n'
front+='''- [Appendix A Repository evidence index](#appendix-a-repository-evidence-index)
- [Appendix B Finding source and verification map](#appendix-b-finding-source-and-verification-map)
- [Appendix C Audit execution record](#appendix-c-audit-execution-record)
- [Appendix D Outstanding confirmations](#appendix-d-outstanding-confirmations)

'''

evidence=(ROOT/'audit_evidence.md').read_text(encoding='utf-8')
entries=[]
for line in evidence.splitlines():
    if re.match(r'\| E\d\d \|',line) and int(line.split('|')[1].strip()[1:]) <= 64:
        cells=[x.strip() for x in line.strip('|').split('|')]
        entries.append(cells)
appA='''# Appendix A Repository evidence index

All factual references below are repository sources. No external publication was newly used as authority, so no external bibliography is added. Tool and library descriptions are based on manifests, imports and configuration, not on claims about their current releases. Historical paths are resolved at the named commit. Current paths are relative to the repository root at the report baseline.

The complete classifications and verification boundaries remain in `audit_evidence.md`. Original prompt material is located under `docs/ai-audit/prompts/`, with its missing exports and edited wording distinguished in Section 3. The four original Mermaid sources remain in `diagrams/`. The report figures are simplified, source-derived visual presentations of the same architecture and are not runtime evidence.

'''
for cells in entries:
    eid,claim,source=cells[:3]
    # Avoid copying long analytical claims already developed in the report.
    typ='Confirmed record' if 'Confirmed' in claim else 'Inferred interpretation'
    appA+=f'**{eid} {typ}.** {source}\n\n'

findings=(ROOT/'audit_findings.md').read_text(encoding='utf-8')
appB='''# Appendix B Finding source and verification map

The severity and status summary is Table 5.1. The following locators preserve the affected functions and validation boundaries for each significant finding. Recommended remedies are not changes applied by this audit. Full finding records remain in `audit_findings.md`.

'''
for m in re.finditer(r'^### ([HCP]\d\d .+?)\n\n(.*?)(?=^### |^## |\Z)',findings,re.M|re.S):
    title,body=m.groups()
    paras=body.strip().split('\n\n')
    selected=[p for p in paras if p.startswith('**Affected') or p.startswith('**Fix and') or p.startswith('**Status and')]
    if title.startswith('P04'):
        selected=[p.replace('Arithmetic reasoning is recorded, but no EVM deployment was executed','The Stage 6 BigInt check confirms the mathematical example, but no EVM deployment was executed') for p in selected]
    appB+='## '+title+'\n\n'+'\n\n'.join(selected)+'\n\n'

testing=(ROOT/'testing_results.md').read_text(encoding='utf-8')
command_table=re.search(r'\| Record \| Exact command.*?(?=\n\n)',testing,re.S).group()
appC='''# Appendix C Audit execution record

The commands below were executed during Stage 6 from the repository root. The environment was Windows PowerShell with Node v25.9.0 and npm 11.12.1. Project dependencies were absent. The audit runner is `validation/stage6_run.ps1`, invoked as `& ./ai_audit_report/validation/stage6_run.ps1`. Log timestamps, commands and exit codes are preserved in the listed files.

Table C.1 Commands and observed results

'''+command_table+'''

The five checks reproduced C01–C05 within their source-level scope. Passing those counterexample assertions does not mean the defects were fixed. The four additional Node tests checked display lifecycle, bounded price decay, estimates for an accepted commitment and P04's mathematical example. No contract compilation, existing Solidity or Mocha suite, coverage, Slither analysis, npm advisory scan, browser session or live transaction was executed during Stage 6. Section 6 explains the historical evidence separately. [E60–E62]

The first Stage 5 check attempt failed in the audit harness while extracting TypeScript for C05. It was corrected without changing application code, and the subsequent five-check run passed. Both attempt logs remain available. The new checks and their hashes are recorded in `testing_results.md` and `validation/`. [E56]

'''
appD='''# Appendix D Outstanding confirmations

The following evidence gaps remain after report integration. They qualify the conclusions rather than being silently filled with assumptions. The complete requests and priorities are preserved in `questions_for_team.md`.

Table D.1 Outstanding evidence and its consequence

| Topic | Open questions | Consequence for this report |
| --- | --- | --- |
| Course requirements and approval | Q01–Q02, Q14 | Primary brief, rubric, lecture material and TA decisions are needed to verify compliance and economic interpretations |
| Prompt and context provenance | Q04–Q05, Q18, Q21–Q22 | Original exports, dated handoffs and compaction or delegation records are missing. Exact prompt wording and context-loss causation remain unverified |
| Human contributions and review | Q06–Q07, Q19 | Member responsibilities and independent peer-review depth cannot be assigned from Git identities or edited logs |
| Reproducible validation | Q08–Q09, Q11, Q15, Q17 | CI artifacts, browser drivers, mutation patches, supported assessment environment and receipts are needed for replay |
| Public deployment | Q10 | An empty deployment map does not establish a Sepolia deployment. Addresses and receipts are required |
| Accepted operational scope | Q12–Q13 | Contract wallets, ownership changes, donations, extremes and reorg handling need explicit expectations |
| Productivity and future practices | Q20, Q23 | No quantified effort saving or adoption of proposed practices is claimed |
| Submission provenance | Q16 | Audit files are uncommitted. The current packaging script requires a clean branch and does not preserve backup-only history |

The stage-outline question Q03 was resolved by the requested eight-stage process. No other team response is assumed. Completion of the report does not close the open functional findings or certify the application for production use.
'''
md=front+'\n\n'.join(parts)+'\n\n'+appA+appB+appC+appD
md=md.replace('\u2014',', ').replace(';',',')
(ROOT/'AI_Audit_Report.md').write_text(md,encoding='utf-8')

# Build DOCX from exactly the integrated Markdown blocks.
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.8)
sec.left_margin=sec.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
    st.font.size=Pt(11)
    st.paragraph_format.space_after=Pt(7)
    st.paragraph_format.line_spacing=1.12
doc.styles['Title'].font.size=Pt(29)
for name,size in [('Heading 1',18),('Heading 2',13),('Heading 3',11)]:
    doc.styles[name].font.size=Pt(size);doc.styles[name].font.bold=True
    doc.styles[name].paragraph_format.space_before=Pt(14)
    doc.styles[name].paragraph_format.keep_with_next=True
doc.styles['Heading 1'].paragraph_format.page_break_before=True
doc.styles['Caption'].font.size=Pt(10)
doc.styles['Caption'].font.italic=True
sec.different_first_page_header_footer=True
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('DutchPoint AI Audit  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for run in footer.runs:run.font.size=Pt(9)
doc.core_properties.title='DutchPoint AI Audit Report'
doc.core_properties.subject='SC4053 Blockchain Technology'
doc.core_properties.author=''
doc.core_properties.comments=''
blocks=[]
def plain(t):
    t=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1',t)
    return t.replace('**','').replace('`','')
def anchor(slug):
    return 'b_'+hashlib.sha256(slug.encode()).hexdigest()[:24]
def add_inline(p,text):
    for piece in re.split(r'(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))',text):
        if not piece:continue
        if piece.startswith('**'):
            r=p.add_run(piece[2:-2]);r.bold=True
        elif piece.startswith('`'):
            r=p.add_run(piece[1:-1]);r.font.name='Consolas';r.font.size=Pt(9)
        elif re.match(r'\[.+\]\(.+\)',piece):
            label,target=re.match(r'\[(.+)\]\((.+)\)',piece).groups()
            if target.startswith('#'):
                h=OxmlElement('w:hyperlink');h.set(qn('w:anchor'),anchor(target[1:]))
                r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)
            else:p.add_run(label)
        else:p.add_run(piece)
def para(text,style=None):
    p=doc.add_paragraph(style=style);add_inline(p,text);blocks.append(plain(text));return p
lines=md.splitlines();i=0;bookmark=1;on_title=True
while i<len(lines):
    line=lines[i]
    if not line.strip():i+=1;continue
    if line.startswith('#'):
        level=len(line)-len(line.lstrip('#'));text=line[level:].strip()
        if text=='DutchPoint AI Audit Report':
            p=para(text,'Title');p.paragraph_format.space_before=Pt(95)
        else:
            if text=='Scope and evidence':doc.add_page_break();on_title=False
            p=para(text,f'Heading {min(level,3)}')
            slug=re.sub(r'[^a-z0-9 -]','',text.lower()).replace(' ','-')
            start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(bookmark));start.set(qn('w:name'),anchor(slug))
            end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(bookmark));bookmark+=1
            p._p.insert(0,start);p._p.append(end)
        i+=1;continue
    if line.startswith('!['):
        alt,path=re.match(r'!\[(.*?)\]\((.*?)\)',line).groups()
        p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(ROOT/path),width=Inches(6.6))
        pic=p._p.xpath('.//wp:docPr')[0];pic.set('descr',alt)
        blocks.append('[IMAGE] '+path);i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            cells=[x.strip() for x in lines[i].strip('|').split('|')]
            if not all(re.fullmatch(r'[: -]+',x) for x in cells):rows.append(cells)
            i+=1
        cols=len(rows[0]);table=doc.add_table(rows=0, cols=cols);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
        weights={2:[.32,.68],3:[.26,.34,.40],4:[.32,.16,.27,.25]}.get(cols,[1/cols]*cols)
        # Keep descriptive cells wide, with a compact exit-code column.
        if rows[0][0]=='Record':weights=[.25,.36,.06,.33]
        for col,width in zip(table.columns,weights):col.width=Inches(6.67*width)
        pr=table._tbl.tblPr
        borders=OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
        pr.append(borders)
        for ri,rowdata in enumerate(rows):
            row=table.add_row();trpr=row._tr.get_or_add_trPr()
            no_split=OxmlElement('w:cantSplit');trpr.append(no_split)
            if ri==0:trpr.append(OxmlElement('w:tblHeader'))
            for ci,(cell,text) in enumerate(zip(row.cells,rowdata)):
                cell.width=Inches(6.67*weights[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr=cell._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'DCE6EF' if ri==0 else ('F6F8FA' if ri%2==0 else 'FFFFFF'));tcpr.append(shd)
                margins=OxmlElement('w:tcMar')
                for edge in ['top','bottom','left','right']:
                    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
                tcpr.append(margins)
                p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.03
                add_inline(p,text)
                for run in p.runs:run.font.size=Pt(9);run.bold=True if ri==0 else run.bold
                if text.isdigit():p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                blocks.append(plain(text))
        doc.add_paragraph().paragraph_format.space_after=Pt(2)
        continue
    if line.startswith('```'):
        i+=1
        while i<len(lines) and not lines[i].startswith('```'):
            p=para(lines[i]);p.paragraph_format.space_after=Pt(2)
            for r in p.runs:r.font.name='Consolas';r.font.size=Pt(9)
            i+=1
        i+=1;continue
    if re.match(r'^(Table|Figure) [A-Z0-9]+\.\d+ [A-Z][^.]*$',line):
        p=para(line,'Caption');p.paragraph_format.keep_with_next=line.startswith('Table');i+=1;continue
    if line.startswith('- '):para(line[2:],'List Bullet');i+=1;continue
    chunk=[line];i+=1
    while i<len(lines) and lines[i].strip() and not re.match(r'^(#|\||!\[|```|- )',lines[i]):chunk.append(lines[i]);i+=1
    p=para(' '.join(chunk))
    if on_title:p.paragraph_format.space_after=Pt(18)

out=ROOT/'AI_Audit_Report.docx';doc.save(out)
(QA/'expected_docx_text.json').write_text(json.dumps(blocks,ensure_ascii=False,indent=2),encoding='utf-8')
after={str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in ROOT.glob('sections/*.md')}
assert before==after
print(json.dumps({'markdown_words':len(md.split()),'tables':len(doc.tables),'figures':len(doc.inline_shapes),'sections_preserved':before==after,'docx_bytes':out.stat().st_size},indent=2))
