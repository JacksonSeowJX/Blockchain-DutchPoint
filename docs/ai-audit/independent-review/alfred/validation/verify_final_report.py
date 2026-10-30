from pathlib import Path
import re,json,hashlib,zipfile
from lxml import etree
from docx import Document

root=Path(__file__).resolve().parents[1]
qa=root/'validation/stage8'
md=(root/'AI_Audit_Report.md').read_text(encoding='utf-8')
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main','r':'http://schemas.openxmlformats.org/package/2006/relationships'}
with zipfile.ZipFile(root/'AI_Audit_Report.docx') as z:
    assert z.testzip() is None
    xml=etree.fromstring(z.read('word/document.xml'))
    for name in z.namelist():
        if name.endswith('.xml') or name.endswith('.rels'):etree.fromstring(z.read(name))
    rel=etree.fromstring(z.read('word/_rels/document.xml.rels'))
    for e in rel:
        if e.get('Type','').endswith('/image'):
            assert 'word/'+e.get('Target') in z.namelist()
    actual=[''.join(p.xpath('.//w:t/text()',namespaces=ns)) for p in xml.xpath('//w:body//w:p',namespaces=ns)]
actual=[x for x in actual if x.strip()]
expected=[x for x in json.loads((qa/'expected_docx_text.json').read_text(encoding='utf-8')) if x.strip() and not x.startswith('[IMAGE]')]
def norm(s):return re.sub(r'\s+',' ',s).strip()
assert list(map(norm,actual))==list(map(norm,expected)), 'DOCX block text differs from canonical integrated content'
md_lines=[]
for line in md.splitlines():
    if not line.strip() or line.startswith('![') or line.startswith('```'):continue
    if line.startswith('|'):
        cells=[c.strip() for c in line.strip('|').split('|')]
        if all(re.fullmatch(r'[: -]+',c) for c in cells):continue
        md_lines.extend(cells)
    else:
        line=re.sub(r'^#{1,3} ','',line)
        line=re.sub(r'^- ','',line)
        md_lines.append(line)
md_text=' '.join(md_lines)
md_text=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1',md_text).replace('**','').replace('`','')
assert norm(md_text)==norm(' '.join(actual)), 'Independent Markdown-to-OOXML text comparison failed'
doc=Document(root/'AI_Audit_Report.docx')
headings=[p.text for p in doc.paragraphs if p.style.name=='Heading 1']
mandatory=re.findall(r'^# ([1-8] .+)$',md,re.M)
assert len(mandatory)==8 and headings[:8]==mandatory
anchors=set(xml.xpath('//w:bookmarkStart/@w:name',namespaces=ns))
links=xml.xpath('//w:hyperlink/@w:anchor',namespaces=ns)
assert all(a in anchors for a in links)
assert len(links)==12
assert len(doc.inline_shapes)==4 and len(doc.tables)==14
assert len(re.findall(r'^Figure 4\.\d [A-Z][^.\n]*$',md,re.M))==4
assert len(re.findall(r'^Table [A-Z0-9]+\.\d [A-Z][^.\n]*$',md,re.M))==14
assert '\u2014' not in md and ';' not in md
assert 'Stage 6 has not begun' not in md
assert '../audit_' not in md
assert set(re.findall(r'\b[HCP]\d{2}\b',md))==set([f'H{i:02}' for i in range(1,7)]+[f'C{i:02}' for i in range(1,6)]+[f'P{i:02}' for i in range(1,5)])
evidence=(root/'audit_evidence.md').read_text(encoding='utf-8')
ids=set(re.findall(r'^\| (E\d+) \|',evidence,re.M))
cited=set(re.findall(r'\bE\d{2}\b',md))
assert cited<=ids
for a,b in re.findall(r'E(\d{2})[–-]E?(\d{2})',md):assert all(f'E{i:02}' in ids for i in range(int(a),int(b)+1))
for path in re.findall(r'!\[[^]]*\]\(([^)]+)\)',md):assert (root/path).exists()
preserved=json.loads((qa/'preserved_sections.json').read_text())
assert all(hashlib.sha256((root/p).read_bytes()).hexdigest()==h for p,h in preserved.items())
result={'package_integrity':'pass','all_xml_well_formed':'pass','python_docx_reopen':'pass','mandatory_sections_in_order':mandatory,'text_blocks_compared':len(actual),'canonical_text_equivalence':'pass','toc_internal_links':len(links),'figures':4,'tables':14,'evidence_ids_resolve':True,'individual_sections_unchanged':True,'rendered_page_review':'NOT COMPLETED - renderer lacks soffice.exe','Word_application_open':'NOT VERIFIED - no Word application available','files':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in [root/'AI_Audit_Report.docx',root/'AI_Audit_Report.md']}}
(qa/'quality_checks.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
print(json.dumps(result,indent=2))
